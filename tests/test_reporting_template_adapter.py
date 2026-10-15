from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src.reporting import build_template_report, summarize_tqs, summarize_utilization


def test_template_adapter_maps_fields_preserves_other_sheets_styles_and_formulas():
    template = Workbook()
    sheet = template.active
    sheet.title = "Allocation"
    sheet.append(["Point Tag", "Physical Slot", "Calculated"])
    sheet.append([None, None, '=A2&"-checked"'])
    sheet["A2"]._style = sheet["A1"]._style
    notes = template.create_sheet("Notes")
    notes["A1"] = "Keep this sheet"
    source = BytesIO()
    template.save(source)
    allocation = pd.DataFrame([
        {"tag": "AI-001", "slot_id": "CAB-1/CH-1"},
        {"tag": "DI-001", "slot_id": "CAB-1/CH-2"},
    ])
    mapping = {"sheet_name": "Allocation", "header_row": 1, "data_start_row": 2,
               "columns": {"Point Tag": "tag", "Physical Slot": "slot_id"}}
    output = build_template_report(allocation, source.getvalue(), mapping)
    result = load_workbook(BytesIO(output), data_only=False)
    assert result.sheetnames == ["Allocation", "Notes"]
    assert result["Allocation"]["A2"].value == "AI-001"
    assert result["Allocation"]["B3"].value == "CAB-1/CH-2"
    assert result["Allocation"]["C2"].value == '=A2&"-checked"'
    assert result["Allocation"]["C3"].value == '=A3&"-checked"'
    assert result["Notes"]["A1"].value == "Keep this sheet"


def test_template_adapter_rejects_missing_headers_and_fields():
    template = Workbook()
    template.active.title = "Allocation"
    template.active.append(["Tag"])
    source = BytesIO()
    template.save(source)
    frame = pd.DataFrame([{"tag": "AI-001"}])
    with pytest.raises(ValueError, match="header 'Slot'"):
        build_template_report(frame, source.getvalue(), {"sheet_name": "Allocation", "columns": {"Slot": "slot_id"}})
    with pytest.raises(ValueError, match="source field 'slot_id'"):
        build_template_report(frame, source.getvalue(), {"sheet_name": "Allocation", "columns": {"Tag": "slot_id"}})


def test_summaries_use_actual_allocations_capacities_and_identical_tq_causes():
    config = {"hardware": {"cabinets": 1, "ciocs_per_cabinet": 1, "baseplates_per_cioc": 1, "charms_per_baseplate": 2}}
    allocation = pd.DataFrame([
        {"tag": "AI-1", "cabinet": "CAB-01", "status": "ALLOCATED"},
        {"tag": "AI-2", "cabinet": "", "status": "TQ"},
    ])
    summary = summarize_utilization(allocation, config)
    assert summary.loc[0, "physical_slots"] == 2
    assert summary.loc[0, "available_slots"] == 2
    assert summary.loc[0, "allocated"] == 1
    assert summary.loc[0, "spare_slots"] == 1
    assert summary.loc[0, "utilization_percent"] == 50
    clusters = summarize_tqs([
        {"tag": "A", "category": "CAPACITY", "rule_id": "r", "reason": "No slot", "suggestion": "Add slots"},
        {"tag": "B", "category": "CAPACITY", "rule_id": "r", "reason": "No slot", "suggestion": "Add slots"},
        {"tag": "C", "category": "CAPACITY", "rule_id": "r", "reason": "Other", "suggestion": "Review"},
    ])
    assert sorted(clusters.affected_count.tolist()) == [1, 2]
