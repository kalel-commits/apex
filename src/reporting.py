"""Workbook generation, template mapping and allocation summaries."""
from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from src.allocator import build_slots


def summarize_utilization(allocation: pd.DataFrame, config: dict[str, Any]) -> pd.DataFrame:
    """Report actual allocated counts against configured physical/available slots."""
    slots = build_slots(config)
    by_cabinet: dict[str, dict[str, int]] = {}
    for slot in slots:
        cabinet = str(slot.get("cabinet_path", slot.get("cabinet", "")))
        counts = by_cabinet.setdefault(cabinet, {"physical_slots": 0, "available_slots": 0, "reserved_slots": 0, "unavailable_slots": 0})
        counts["physical_slots"] += 1
        availability = slot.get("availability", "available")
        if availability == "available":
            counts["available_slots"] += 1
        elif availability == "reserved":
            counts["reserved_slots"] += 1
        else:
            counts["unavailable_slots"] += 1

    allocated = allocation[allocation.get("status", pd.Series(index=allocation.index, dtype=str)).eq("ALLOCATED")]
    used = allocated.groupby("cabinet", dropna=False).size().to_dict() if "cabinet" in allocated else {}
    rows = []
    for cabinet, capacities in sorted(by_cabinet.items()):
        count = int(used.get(cabinet, 0))
        available = capacities["available_slots"]
        rows.append({
            "cabinet": cabinet,
            **capacities,
            "allocated": count,
            "spare_slots": available - count,
            "utilization_percent": round(100 * count / available, 2) if available else 0.0,
        })
    return pd.DataFrame(rows, columns=["cabinet", "physical_slots", "available_slots", "reserved_slots", "unavailable_slots", "allocated", "spare_slots", "utilization_percent"])


def summarize_tqs(tqs) -> pd.DataFrame:
    """Cluster only identical emitted causes; clusters are not inferred root causes."""
    columns = ["remediation_rank", "category", "rule_id", "reason", "suggestion", "affected_count", "signal_tags"]
    frame = pd.DataFrame(tqs)
    if frame.empty:
        return pd.DataFrame(columns=columns)
    for name in ("category", "rule_id", "reason", "suggestion", "tag"):
        if name not in frame:
            frame[name] = ""
    grouped = frame.groupby(["category", "rule_id", "reason", "suggestion"], dropna=False, sort=True)
    rows = []
    for (category, rule_id, reason, suggestion), group in grouped:
        tags = sorted(set(str(value) for value in group["tag"] if str(value)))
        rows.append({"category": category, "rule_id": rule_id, "reason": reason, "suggestion": suggestion,
                     "affected_count": len(tags), "signal_tags": ", ".join(tags)})
    rows.sort(key=lambda row: (-row["affected_count"], str(row["category"]), str(row["rule_id"]), str(row["reason"])))
    for rank, row in enumerate(rows, 1):
        row["remediation_rank"] = rank
    return pd.DataFrame(rows, columns=columns)


def _write_table_sheet(writer, name: str, frame: pd.DataFrame) -> None:
    frame.to_excel(writer, index=False, sheet_name=name)


def _report_frames(allocation, tqs, validation, metadata=None, revision=None, config=None):
    metadata = dict(metadata or allocation.attrs.get("metadata", {}))
    alloc = allocation.copy()
    alloc.attrs.clear()
    is_allocated = alloc["status"].eq("ALLOCATED") if "status" in alloc else pd.Series(False, index=alloc.index)
    allocated = alloc[is_allocated].copy()
    unallocated = alloc[~is_allocated].copy()
    utilization = summarize_utilization(alloc, config) if config else pd.DataFrame(columns=[
        "cabinet", "physical_slots", "available_slots", "reserved_slots", "unavailable_slots", "allocated", "spare_slots", "utilization_percent"
    ])
    tq_summary = summarize_tqs(tqs)
    summary = {
        "allocation_status": metadata.get("status", "UNKNOWN"),
        "validation_status": "IMPLEMENTED CHECKS PASS; ENGINEERING VALIDITY UNVERIFIED" if not validation and not tqs and len(alloc) and is_allocated.all() else "INVALID OR INCOMPLETE",
        "total_signals": len(alloc), "allocated": len(allocated), "unallocated": len(unallocated),
        "unresolved_tqs": len(tqs), "validation_findings": len(validation),
        "cabinet_count_used": int((utilization.allocated > 0).sum()) if not utilization.empty else 0,
        "cabinet_count_configured": len(utilization),
        "inventory_authority": (config or {}).get("hardware", {}).get("inventory", {}).get("source_authority", "synthetic or incomplete"),
        "all_mandatory_requirements_verified": False,
    }
    if revision is not None:
        rev = pd.DataFrame(revision)
        counts = rev.revision_status.value_counts() if "revision_status" in rev else {}
        summary.update({f"revision_{name.lower()}": int(counts.get(name, 0)) for name in ("ADDED", "REMOVED", "MODIFIED", "MOVED", "UNCHANGED", "UNALLOCATED")})
        if "weighted_move_cost_points" in rev:
            summary["lifecycle_weighted_move_cost_points"] = float(rev.weighted_move_cost_points.sum())
    return alloc, allocated, unallocated, utilization, tq_summary, metadata, summary


def build_excel_report(allocation: pd.DataFrame, tqs, validation, metadata=None, revision=None, config=None) -> bytes:
    """Build a new standard workbook and reconcile its allocation rows on reopen."""
    alloc, allocated, unallocated, utilization, tq_summary, metadata, summary = _report_frames(
        allocation, tqs, validation, metadata, revision, config
    )
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _write_table_sheet(writer, "IO Allocation", alloc)
        _write_table_sheet(writer, "Unallocated Signals", unallocated)
        _write_table_sheet(writer, "Technical Queries", pd.DataFrame(tqs))
        _write_table_sheet(writer, "TQ Summary", tq_summary)
        _write_table_sheet(writer, "Validation", pd.DataFrame(validation))
        _write_table_sheet(writer, "Utilization", utilization)
        _write_table_sheet(writer, "Run Metadata", pd.DataFrame([{"metric": key, "value": value} for key, value in metadata.items()]))
        _write_table_sheet(writer, "Summary", pd.DataFrame([summary]))
        if revision is not None:
            _write_table_sheet(writer, "Revision Comparison", pd.DataFrame(revision))

    content = buffer.getvalue()
    check = pd.ExcelFile(BytesIO(content), engine="openpyxl")
    expected = {"IO Allocation", "Technical Queries", "TQ Summary", "Validation", "Unallocated Signals", "Utilization", "Run Metadata", "Summary"}
    if not expected.issubset(check.sheet_names):
        raise ValueError("Generated workbook is missing required worksheets.")
    reread = pd.read_excel(BytesIO(content), sheet_name="IO Allocation", engine="openpyxl")
    if len(reread) != len(alloc) or list(reread.columns) != list(alloc.columns):
        raise ValueError("Generated workbook failed allocation row/column verification.")
    if "tag" in reread and reread.tag.fillna("").astype(str).tolist() != alloc.tag.fillna("").astype(str).tolist():
        raise ValueError("Generated workbook failed signal tag reconciliation.")
    if len(pd.read_excel(BytesIO(content), sheet_name="Unallocated Signals")) != len(unallocated):
        raise ValueError("Generated workbook failed allocated/unallocated reconciliation.")
    return content


def build_template_report(allocation: pd.DataFrame, template, mapping: dict[str, Any]) -> bytes:
    """Fill configured allocation columns in a supplied workbook, preserving its sheets/styles.

    `mapping` requires `sheet_name`, `columns` (template header -> allocation field),
    and optionally `header_row` / `data_start_row`. This cannot establish official
    template compatibility until exercised against the organizer's actual workbook.
    """
    if not isinstance(mapping, dict) or not isinstance(mapping.get("columns"), dict) or not mapping["columns"]:
        raise ValueError("Template mapping must define a nonempty 'columns' object.")
    sheet_name = mapping.get("sheet_name")
    if not isinstance(sheet_name, str) or not sheet_name:
        raise ValueError("Template mapping requires a worksheet 'sheet_name'.")
    header_row = mapping.get("header_row", 1)
    data_start = mapping.get("data_start_row", header_row + 1)
    if not isinstance(header_row, int) or header_row < 1 or not isinstance(data_start, int) or data_start <= header_row:
        raise ValueError("Template header_row/data_start_row must be positive ordered row numbers.")
    content = template.getvalue() if hasattr(template, "getvalue") else template
    workbook = load_workbook(BytesIO(content))
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Template worksheet {sheet_name!r} was not found.")
    sheet = workbook[sheet_name]
    headers = {str(cell.value).strip().casefold(): cell.column for cell in sheet[header_row] if cell.value is not None}
    column_map = {}
    for header, source in mapping["columns"].items():
        if str(header).strip().casefold() not in headers:
            raise ValueError(f"Template header {header!r} was not found on worksheet {sheet_name!r}.")
        if source not in allocation.columns:
            raise ValueError(f"Allocation source field {source!r} was not found.")
        column_map[headers[str(header).strip().casefold()]] = source

    max_clear = max(sheet.max_row, data_start + len(allocation) - 1)
    for column in column_map:
        for row_number in range(data_start, max_clear + 1):
            sheet.cell(row_number, column).value = None

    style_row = data_start
    for offset, (_, record) in enumerate(allocation.iterrows()):
        row_number = data_start + offset
        if row_number != style_row:
            for column in range(1, sheet.max_column + 1):
                source_cell = sheet.cell(style_row, column)
                target_cell = sheet.cell(row_number, column)
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
                if source_cell.number_format:
                    target_cell.number_format = source_cell.number_format
                if source_cell.data_type == "f" and column not in column_map:
                    target_cell.value = Translator(source_cell.value, origin=source_cell.coordinate).translate_formula(target_cell.coordinate)
        for column, source in column_map.items():
            sheet.cell(row_number, column).value = record[source]

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
