"""Input parsing and conservative column normalization."""
from __future__ import annotations
import re
import csv
import io
from typing import BinaryIO, Mapping
from pathlib import Path
import pandas as pd

ALIASES: dict[str, set[str]] = {
    "tag": {"tag", "signal", "signal tag", "instrument tag", "io tag", "point tag", "tag name"},
    "io_type": {"io", "i/o", "io type", "signal type", "type"},
    "group": {"grp", "group", "signal group", "functional group"},
    "segment": {"segment", "io segment"},
    "redundancy_group": {"redundancy group", "redundancy_group", "ab group", "ab_group"},
    "redundancy_side": {"redundancy side", "redundancy_side", "side", "ab side"},
    "duty_standby": {"duty/standby", "duty standby", "duty_standby", "service role"},
    "cable": {"cable", "cable no", "cable number", "cable id"},
    "pair": {"pair", "pair no", "pair number"},
    "charm_type": {"charm type", "required charm", "charm"},
    "terminal_block_type": {"terminal block type", "tb type", "terminal block"},
    "isolation_required": {"isolation required", "requires isolation", "isolation requirement"},
    "fusing_required": {"fusing required", "fuse required", "requires fusing"},
    "injected_power_required": {"injected power required", "requires injected power", "power injection"},
    "description": {"description", "signal description", "service description"},
    "existing_cabinet": {"existing cabinet", "current cabinet", "old cabinet"},
    "previous_location": {"previous location", "previous_location", "old slot", "rev a location"},
    "lifecycle": {"lifecycle", "engineering status"},
}
REQUIRED = {"tag"}

class _NamedStream:
    def __init__(self, stream, name): self._stream,self.name=stream,name
    def __getattr__(self, attribute): return getattr(self._stream,attribute)

def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).strip().lower()).strip()

def _check_source_headers(headers, mapping=None):
    aliases = {k:{norm(x) for x in v} for k,v in ALIASES.items()}
    explicit=dict(mapping or {}); claimed={}; raw_seen={}
    for header in headers:
        if header is None or not str(header).strip(): continue
        raw=norm(header)
        if raw in raw_seen: raise ValueError(f"Duplicate input header {header!r} (also {raw_seen[raw]!r}).")
        raw_seen[raw]=header
        target=explicit.get(str(header))
        if target is None:
            matches=[field for field,values in aliases.items() if raw in values]
            if len(matches)>1: raise ValueError(f"Ambiguous header mapping for {header!r}: {matches}")
            target=matches[0] if matches else None
        if target is not None:
            if target in claimed: raise ValueError(f"Ambiguous mapping: both {claimed[target]!r} and {header!r} map to {target!r}.")
            claimed[target]=header

def normalize_columns(df: pd.DataFrame, mapping: Mapping[str, str] | None = None,
                      aliases: Mapping[str, list[str] | set[str]] | None = None) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame): raise TypeError("Input must be a table.")
    if df.columns.duplicated().any():
        raise ValueError(f"Duplicate input headers: {df.columns[df.columns.duplicated()].tolist()}")
    alias_map = {k: {norm(x) for x in v} for k, v in (aliases or ALIASES).items()}
    explicit = dict(mapping or {})  # source header -> canonical field
    rename: dict[object, str] = {}
    claimed: dict[str, object] = {}
    for col in df.columns:
        key = explicit.get(str(col))
        matches = [field for field, values in alias_map.items() if norm(col) in values]
        if key is None:
            if len(matches) > 1: raise ValueError(f"Ambiguous header mapping for {col!r}: {matches}")
            key = matches[0] if matches else None
        if key:
            if key not in alias_map: raise ValueError(f"Unsupported mapped field {key!r} for {col!r}.")
            if key in claimed: raise ValueError(f"Ambiguous mapping: both {claimed[key]!r} and {col!r} map to {key!r}.")
            claimed[key] = col; rename[col] = key
    result = df.rename(columns=rename).copy()
    missing = REQUIRED - set(result.columns)
    if missing: raise ValueError(f"Missing required fields: {sorted(missing)}. Map an input column explicitly.")
    # Entirely blank spreadsheet rows are noise; partially blank records are not.
    result = result.dropna(how="all").reset_index(drop=True)
    result["tag"] = result["tag"].fillna("").astype(str).str.strip()
    if result.tag.eq("").any(): raise ValueError(f"Missing signal tags on rows {(result.index[result.tag.eq('')]+2).tolist()[:10]}.")
    if result.tag.duplicated().any():
        raise ValueError(f"Duplicate signal tags: {result.loc[result.tag.duplicated(keep=False), 'tag'].tolist()[:10]}")
    for col in ALIASES:
        if col not in result: result[col] = ""
        result[col] = result[col].fillna("").astype(str).str.strip()
    return result

def load_input(file_obj: BinaryIO, mapping: Mapping[str, str] | None = None, sheet_name: str | int = 0) -> pd.DataFrame:
    if isinstance(file_obj,(str,Path)):
        path=Path(file_obj)
        with path.open("rb") as stream:
            return load_input(_NamedStream(stream,str(path)),mapping=mapping,sheet_name=sheet_name)
    name = str(getattr(file_obj, "name", "input.xlsx")).lower()
    try:
        start=file_obj.tell()
        if name.endswith(".csv"):
            wrapper=io.TextIOWrapper(file_obj,encoding="utf-8-sig",newline="")
            try: headers=next(csv.reader(wrapper),[])
            finally: wrapper.detach(); file_obj.seek(start)
            _check_source_headers(headers,mapping)
            frame = pd.read_csv(file_obj,encoding="utf-8-sig")
        elif name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            workbook=load_workbook(file_obj,read_only=True,data_only=True)
            try:
                worksheet=workbook.worksheets[sheet_name] if isinstance(sheet_name,int) else workbook[sheet_name]
                _check_source_headers(next(worksheet.iter_rows(min_row=1,max_row=1,values_only=True),()),mapping)
            finally:
                workbook.close(); file_obj.seek(start)
            frame = pd.read_excel(file_obj, sheet_name=sheet_name)
        else: raise ValueError("Supported input formats are .csv, .xlsx, and .xlsm.")
    except Exception as exc:
        if isinstance(exc, ValueError) and str(exc).startswith(("Supported","Duplicate input header","Ambiguous mapping","Ambiguous header mapping")): raise
        raise ValueError(f"Could not read input workbook: {exc}") from exc
    return normalize_columns(frame, mapping=mapping)
